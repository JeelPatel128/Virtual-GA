"""Knowledge base: turn every upload into text, chunk it, embed it, and search it."""
import csv
import io
import math
import re
from collections import Counter

import numpy as np
from sqlalchemy import delete, select

from . import llm
from .config import BASE_DIR
from .db import Chunk, Doc

CHUNK_CHARS = 900
OVERLAP = 150

STOPWORDS = set("""a an the is are was were be been to of in on at for and or but if with by from as it its
this that these those i me my we our you your he she they them do does did can could how what when where which
who whom why will would should there here about into over under than then so not no yes please hi hello""".split())


# ---------- extraction ----------

def extract_text(filename: str, data: bytes, mime: str) -> tuple[str, str]:
    """Returns (kind, text) for documents. Images and video are handled separately."""
    name = filename.lower()
    if name.endswith(".pdf") or mime == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "pdf", "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "docx", "\n".join(parts)
    if name.endswith((".txt", ".md", ".csv")) or mime.startswith("text/"):
        return "text", data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {filename}")


def chunk_text(text: str) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= CHUNK_CHARS:
        return [text] if text else []
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        while len(p) > CHUNK_CHARS:  # very long paragraph: hard split
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(p[:CHUNK_CHARS])
            p = p[CHUNK_CHARS - OVERLAP:]
        if len(cur) + len(p) + 2 > CHUNK_CHARS and cur:
            chunks.append(cur)
            cur = cur[-OVERLAP:] + "\n\n" + p
        else:
            cur = f"{cur}\n\n{p}" if cur else p
    if cur:
        chunks.append(cur)
    return chunks


def add_document(db, title: str, kind: str, text: str, added_by: str | None = None,
                 image_bytes: bytes | None = None, mime: str | None = None) -> Doc:
    text = text.strip()
    if not text:
        raise ValueError("Nothing readable was found in that upload.")
    doc = Doc(title=title.strip()[:300], kind=kind, extracted_text=text, image_bytes=image_bytes,
              mime=mime, added_by=added_by)
    db.add(doc)
    db.flush()
    # Prefix every chunk with the title so retrieval knows what it belongs to.
    pieces = chunk_text(text)
    vectors = llm.embed([f"{title}\n{p}" for p in pieces]) or [None] * len(pieces)
    for p, v in zip(pieces, vectors):
        db.add(Chunk(doc_id=doc.id, text=p, embedding=v))
    db.commit()
    return doc


def replace_chunks(db, doc: Doc, text: str, search_prefix: str | None = None) -> None:
    """Deletes a Doc's existing chunks and re-chunks/re-embeds from scratch - used
    whenever an existing Doc's content changes in place (FAQ re-import, a Q&A edit)."""
    db.execute(delete(Chunk).where(Chunk.doc_id == doc.id))
    pieces = chunk_text(text)
    prefix = search_prefix if search_prefix is not None else doc.title
    vectors = llm.embed([f"{prefix}\n{p}" for p in pieces]) or [None] * len(pieces)
    for p, v in zip(pieces, vectors):
        db.add(Chunk(doc_id=doc.id, text=p, embedding=v))


def update_qa_doc(db, doc_id: int, question: str, answer: str, added_by: str | None = None) -> Doc:
    doc = db.get(Doc, doc_id)
    if not doc or doc.kind != "qa":
        raise ValueError("Q&A not found")
    text = f"Q: {question.strip()}\nA: {answer.strip()}"
    doc.title = question.strip()[:300]
    doc.extracted_text = text
    if added_by:
        doc.added_by = added_by
    db.flush()
    replace_chunks(db, doc, text)
    db.commit()
    return doc


# ---------- FAQ bulk import (data/virtual-ga-faq.csv) ----------

FAQ_CSV_COLUMNS = ["id", "category", "type", "question", "variants", "asked_by", "times_mentioned",
                   "answer", "source", "handled_by", "changes_each_semester", "status"]
FAQ_SKIP_STATUS = "Needs answer"


def parse_faq_csv(text: str) -> dict:
    """Matches columns by header name, not position, so a reordered file still imports.
    Returns {"missing_columns": [...], "rows": []} if any expected column is absent -
    the caller shows that list as a clear error rather than guessing."""
    reader = csv.DictReader(io.StringIO(text))
    header = reader.fieldnames or []
    missing = [c for c in FAQ_CSV_COLUMNS if c not in header]
    if missing:
        return {"missing_columns": missing, "rows": []}
    rows = []
    for raw in reader:
        get = lambda k: (raw.get(k) or "").strip()
        status = get("status")
        times = get("times_mentioned")
        rows.append({
            "id": get("id"), "category": get("category"), "type": get("type"),
            "question": get("question"),
            "variants": [v.strip() for v in get("variants").split("|") if v.strip()],
            "asked_by": get("asked_by"),
            "times_mentioned": int(times) if times.isdigit() else None,
            "answer": get("answer"), "source": get("source"), "handled_by": get("handled_by"),
            "changes_each_semester": get("changes_each_semester").lower() == "yes",
            "status": status, "will_import": status != FAQ_SKIP_STATUS,
        })
    return {"missing_columns": [], "rows": rows}


def upsert_faq_doc(db, row: dict, added_by: str | None = None) -> Doc:
    """Inserts or updates (by external_id, i.e. the CSV's own id column) a Q&A Doc from
    one imported FAQ row. Variants (alternate phrasings) are folded into the text used
    for embedding/keyword search, but kept out of the displayed title/answer."""
    question, answer = row["question"], row["answer"]
    text = f"Q: {question}\nA: {answer}"
    doc = db.scalars(select(Doc).where(Doc.external_id == row["id"])).first() if row["id"] else None
    if doc is None:
        doc = Doc(kind="qa", external_id=row["id"] or None)
        db.add(doc)
    doc.title = question[:300]
    doc.extracted_text = text
    doc.category = row["category"] or None
    doc.faq_type = row["type"] or None
    doc.variants = row["variants"]
    doc.asked_by = row["asked_by"] or None
    doc.times_mentioned = row["times_mentioned"]
    doc.source = row["source"] or None
    doc.handled_by = row["handled_by"] or None
    doc.changes_each_semester = row["changes_each_semester"]
    doc.status = row["status"] or None
    if added_by:
        doc.added_by = added_by
    db.flush()
    search_prefix = "\n".join([question, *row["variants"]])
    replace_chunks(db, doc, text, search_prefix)
    db.commit()
    return doc


def import_faq_csv(db, text: str, added_by: str | None = None) -> dict:
    parsed = parse_faq_csv(text)
    if parsed["missing_columns"]:
        return {"missing_columns": parsed["missing_columns"], "imported": 0, "skipped": 0, "errors": []}
    imported = skipped = 0
    errors = []
    for row in parsed["rows"]:
        if not row["will_import"]:
            skipped += 1
            continue
        try:
            upsert_faq_doc(db, row, added_by)
            imported += 1
        except Exception as e:
            errors.append(f"{row['id'] or '(no id)'}: {e}")
    return {"missing_columns": [], "imported": imported, "skipped": skipped, "errors": errors}


# ---------- GA resources library (data/ga_resources.csv) ----------

RESOURCES_CSV_COLUMNS = ["resource_id", "title", "source_file", "audience", "faq_ids",
                         "summary", "seasonal", "notes"]
RESOURCES_DIR = BASE_DIR / "data" / "resources"


def parse_resources_csv(text: str) -> dict:
    reader = csv.DictReader(io.StringIO(text))
    header = reader.fieldnames or []
    missing = [c for c in RESOURCES_CSV_COLUMNS if c not in header]
    if missing:
        return {"missing_columns": missing, "rows": []}
    rows = []
    for raw in reader:
        get = lambda k: (raw.get(k) or "").strip()
        rows.append({
            "resource_id": get("resource_id"), "title": get("title"), "source_file": get("source_file"),
            "audience": get("audience") or "public",
            "faq_ids": [f.strip() for f in get("faq_ids").split(",") if f.strip()],
            "summary": get("summary"), "seasonal": get("seasonal").lower() == "yes",
        })
    return {"missing_columns": [], "rows": rows}


def _resource_text(source_file: str, summary: str) -> tuple[str, str]:
    """(kind, text) - the file's own extracted text if it's been placed under
    data/resources/, falling back to just the CSV's summary otherwise (e.g. before the
    real files exist yet). Resource PDFs are committed to the repo, unlike the gitignored
    directory/inventory CSVs, so reading them here (admin-triggered import, not a live
    visitor request) doesn't cross the "never read data/ at request time" line."""
    path = RESOURCES_DIR / source_file
    if path.is_file():
        try:
            kind, text = extract_text(path.name, path.read_bytes(), "application/octet-stream")
            return kind, (f"{summary}\n\n{text}" if summary else text).strip()
        except ValueError:
            pass
    return "text", summary.strip()


def upsert_resource_doc(db, row: dict) -> Doc:
    kind, text = _resource_text(row["source_file"], row["summary"])
    if not text:
        raise ValueError(f"no summary, and no file found at data/resources/{row['source_file']}")
    doc = db.scalars(select(Doc).where(Doc.external_id == row["resource_id"])).first()
    if doc is None:
        doc = Doc(external_id=row["resource_id"])
        db.add(doc)
    doc.title, doc.kind, doc.extracted_text = row["title"][:300], kind, text
    doc.audience, doc.seasonal, doc.linked_faq_ids = row["audience"], row["seasonal"], row["faq_ids"]
    doc.source_file = row["source_file"] or None
    db.flush()
    replace_chunks(db, doc, text)
    db.commit()
    return doc


def import_resources_csv(db, text: str) -> dict:
    parsed = parse_resources_csv(text)
    if parsed["missing_columns"]:
        return {"missing_columns": parsed["missing_columns"], "imported": 0, "errors": []}
    imported = 0
    errors = []
    for row in parsed["rows"]:
        if not row["resource_id"] or not row["title"]:
            continue
        try:
            upsert_resource_doc(db, row)
            imported += 1
        except Exception as e:
            errors.append(f"{row['resource_id']}: {e}")
    return {"missing_columns": [], "imported": imported, "errors": errors}


def reembed_missing(db) -> int:
    """Embed chunks that were added before the Gemini key existed."""
    rows = db.execute(select(Chunk, Doc.title).join(Doc).where(Chunk.embedding.is_(None))).all()
    if not rows or not llm.enabled():
        return 0
    vectors = llm.embed([f"{t}\n{c.text}" for c, t in rows])
    if not vectors:
        return 0
    for (c, _), v in zip(rows, vectors):
        c.embedding = v
    db.commit()
    return len(rows)


# ---------- retrieval ----------

def _tokens(s: str) -> list[str]:
    return [w for w in re.findall(r"[a-z0-9]+", s.lower()) if w not in STOPWORDS and len(w) > 1]


def search(db, query: str, k: int = 5, visitor: bool = True) -> list[dict]:
    """visitor=True (the default, used by the bot) only ever includes Doc rows marked
    audience="public" - an allowlist, not a denylist, so "ga_only", "approval_required"
    and any future restricted audience value are excluded from visitor-facing search
    without needing to update this check. A seasonal=True doc is excluded the same way.
    The admin "Test what it finds" tool and the GA console's own resource search both
    pass visitor=False to see everything."""
    rows = db.execute(select(Chunk, Doc).join(Doc)).all()
    if visitor:
        rows = [(c, d) for c, d in rows if d.audience == "public" and not d.seasonal]
    if not rows:
        return []
    results = []
    qv = llm.embed([query], for_query=True)
    if qv and all(c.embedding for c, _ in rows):
        q = np.array(qv[0])
        q = q / (np.linalg.norm(q) or 1)
        for c, d in rows:
            v = np.array(c.embedding)
            score = float(q @ (v / (np.linalg.norm(v) or 1)))
            results.append((score, c, d))
        method = "vector"
    else:
        # Keyword fallback (BM25-lite) so the bot still works without an API key.
        # Score = share of the query's (idf-weighted) words found in the chunk, 0 to 1.
        qt = set(_tokens(query))
        # FAQ variants (alternate phrasings) are searchable even though they're kept out
        # of the displayed title/chunk text.
        docs_tokens = [set(_tokens(d.title + " " + c.text + " " + " ".join(d.variants or []))) for c, d in rows]
        n = len(rows)
        df = Counter(t for dt in docs_tokens for t in dt)
        idf = {t: math.log(1 + n / (df[t] or 0.5)) for t in qt}
        total = sum(idf.values()) or 1
        for (c, d), dt in zip(rows, docs_tokens):
            results.append((sum(idf[t] for t in qt if t in dt) / total, c, d))
        method = "keyword"
    results.sort(key=lambda r: r[0], reverse=True)
    return [
        {"id": c.id, "doc_id": d.id, "title": d.title, "kind": d.kind, "text": c.text,
         "score": round(s, 3), "method": method, "faq_type": d.faq_type, "category": d.category}
        for s, c, d in results[:k] if s > 0
    ]
